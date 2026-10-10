from pathlib import Path
import json, re, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent
DATA = OUT / 'guide_content.json'
SOFTWARE_VERSION, GUIDE_EDITION = '8.0.1', '1.1'   # the release the guide describes; the guide's own edition (also stated on the cover page in guide_content.json)

def field(p, instruction):
    r=p.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); r._r.append(begin)
    r=p.add_run(); text=OxmlElement('w:instrText'); text.set(qn('xml:space'),'preserve'); text.text=' '+instruction+' '; r._r.append(text)
    r=p.add_run(); end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); r._r.append(end)

def link(p, label, url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    c=OxmlElement('w:color'); c.set(qn('w:val'),'185B64'); props.append(c)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); props.append(u)
    run.append(props); t=OxmlElement('w:t'); t.text=label; run.append(t); h.append(run); p._p.append(h)

def set_language(style, lang):
    pr=style.element.get_or_add_rPr(); x=OxmlElement('w:lang'); x.set(qn('w:val'),lang); pr.append(x)

def paragraph(doc, text, style=None):
    p=doc.add_paragraph(style=style)
    # Backticks denote literal UI labels, paths and commands; ** is restrained emphasis.
    for part in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)',text):
        if not part: continue
        r=p.add_run(part.strip('`') if part.startswith('`') else part.strip('*') if part.startswith('**') else part)
        if part.startswith('`'): r.font.name='Menlo'; r.font.size=Pt(9.2)
        if part.startswith('**'): r.bold=True
    return p

def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths): c.width=Cm(w)
    for c,h in zip(t.rows[0].cells,headers):
        c.text=h
        for r in c.paragraphs[0].runs:r.bold=True
    rep=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        cells=t.add_row().cells
        for c,text in zip(cells,row):
            c.text=''; p=c.paragraphs[0]
            for part in re.split(r'(`[^`]+`)',text):
                r=p.add_run(part.strip('`') if part.startswith('`') else part)
                if part.startswith('`'):r.font.name='Menlo';r.font.size=Pt(8.5)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        no=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no)
        for ci,c in enumerate(row.cells):
            from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:c.width=Cm(widths[ci])
            pr=c._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
            for edge,value in [('top','90'),('bottom','90'),('left','105'),('right','105')]:
                x=OxmlElement('w:'+edge);x.set(qn('w:w'),value);x.set(qn('w:type'),'dxa');margins.append(x)
            pr.append(margins)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F2F3');pr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.08
                for r in p.runs:
                    if r.font.size is None:r.font.size=Pt(9.2)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def build(lang, content):
    doc=Document();sec=doc.sections[0]
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.65)
    sec.left_margin=Cm(1.85);sec.right_margin=Cm(1.85)
    sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
    sec.different_first_page_header_footer=True
    for name in ['Normal','Body Text','List Bullet','List Number']:
        s=doc.styles[name];s.font.name='Calibri';s.font.size=Pt(10.5);s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.line_spacing=1.14;s.paragraph_format.space_after=Pt(6)
        set_language(s, 'fr-CA' if lang=='FR' else 'en-CA')
    for name,size in [('Title',36),('Subtitle',18),('Heading 1',22),('Heading 2',13),('Heading 3',11)]:
        s=doc.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
        s.font.bold=name!='Subtitle';s.paragraph_format.space_before=Pt(9);s.paragraph_format.space_after=Pt(7)
        s.paragraph_format.keep_with_next=True;set_language(s,'fr-CA' if lang=='FR' else 'en-CA')
    if 'Code' not in doc.styles:
        from docx.enum.style import WD_STYLE_TYPE
        s=doc.styles.add_style('Code',WD_STYLE_TYPE.PARAGRAPH)
    s=doc.styles['Code'];s.font.name='Menlo';s.font.size=Pt(8.6);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.line_spacing=1.06;s.paragraph_format.space_after=Pt(7)
    s.paragraph_format.keep_together=True;s.paragraph_format.left_indent=Cm(.18)
    s.paragraph_format.right_indent=Cm(.05)
    hp=sec.header.paragraphs[0];hp.text='YUCLAW 8  |  '+('Guide utilisateur' if lang=='FR' else 'User Guide')
    hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor(0,0,0)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=fp.add_run((f'YUCLAW {SOFTWARE_VERSION}  •  Édition du guide {GUIDE_EDITION}  •  ' if lang=='FR' else f'YUCLAW {SOFTWARE_VERSION}  •  Guide edition {GUIDE_EDITION}  •  '));r.font.size=Pt(8)
    field(fp,'PAGE')
    doc.core_properties.title='YUCLAW 8 '+('Guide utilisateur français' if lang=='FR' else 'English User Guide')
    doc.core_properties.subject='YUCLAW 8.0.1 local research workbench and public research interface'
    doc.core_properties.author='YUCLAW'
    doc.core_properties.keywords='YUCLAW, 8.0.1, user guide, SHD, EVO, COM, PRC'
    for i,page in enumerate(content):
        if i:doc.add_page_break()
        if page.get('cover'):
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(52)
            paragraph(doc,'YUCLAW', 'Title')
            paragraph(doc,page['title'],'Title')
            paragraph(doc,page['subtitle'],'Subtitle')
        else:
            paragraph(doc,page['title'],'Heading 1')
        for block in page['blocks']:
            k=block[0]
            if k=='p':paragraph(doc,block[1])
            elif k=='h':paragraph(doc,block[1],'Heading 2')
            elif k=='b':paragraph(doc,block[1],'List Bullet')
            elif k=='n':paragraph(doc,block[1],'List Number')
            elif k=='code':
                p=doc.add_paragraph(style='Code');p.add_run(block[1])
            elif k=='table':table(doc,block[1],block[2],block[3] if len(block)>3 else None)
            elif k=='link':
                p=doc.add_paragraph();link(p,block[1],block[2]);
                if len(block)>3:p.add_run(' — '+block[3])
            elif k=='gap':doc.add_paragraph().paragraph_format.space_after=Pt(block[1])
    path=OUT/f'YUCLAW-v8-User-Guide-{lang}.docx';doc.save(path)
    return path

if __name__=='__main__':
    data=json.loads(DATA.read_text())
    for lang,content in data.items():
        path=build(lang,content);print(path)
    en=[b[1] for p in data['EN'] for b in p['blocks'] if b[0]=='code']
    fr=[b[1] for p in data['FR'] for b in p['blocks'] if b[0]=='code']
    print('EN/FR code blocks identical:',en==fr,'blocks:',len(en))
    assert en==fr
